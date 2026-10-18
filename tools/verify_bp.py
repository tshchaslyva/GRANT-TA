# -*- coding: utf-8 -*-
"""
Незалежна перевірка відкоригованого бізнес-плану (після перерахунку в LibreOffice/Excel).
Запуск: python3 tools/verify_bp.py <перерахований_xlsx> <шаблон_xlsx>
Перевіряє: ліміти п. 7 Порядку, узгодженість сум, покриття гранту податками, відсутність дефіциту коштів,
кадрову умову, збереження формул шаблону, відповідність текстових цифр фінансовій таблиці.
"""
import sys, os, re
import openpyxl
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.dirname(__file__))
import model as M

path, template = sys.argv[1], sys.argv[2]
wbv = openpyxl.load_workbook(path, data_only=True)
wbf = openpyxl.load_workbook(path, data_only=False)
wbt = openpyxl.load_workbook(template, data_only=False)
bp, pl = wbv["Бізнес план"], wbv["Прибутки та збитки"]
ok = True
def check(cond, msg):
    global ok
    print(("OK   " if cond else "FAIL ") + msg)
    ok = ok and bool(cond)

grant = bp["C26"].value
check(grant == M.GRANT == pl["C10"].value == bp["E89"].value, f"сума гранту узгоджена: C26={grant}, P&L C10={pl['C10'].value}, інвестплан E89={bp['E89'].value}")
check(500_000 <= grant <= 1_500_000, "грант у базовому діапазоні масштабування 500 000–1 500 000")
check(bp["B24"].value == "Масштабування власної справи" and bp["B22"].value == "36 місяців", "тип проєкту і строк")
check(all(bp[c].value == "Ні" for c in ("C30","C31","C32","C33","C34","C35","C36","C37")) and bp["C29"].value == "Неплатник ПДВ", "підстави для збільшення гранту не заявлені")
check(bp["D39"].value == "Сума в межах Порядку" and bp["E39"].value == "1+ нове місце заплановано", f"контролі шаблону: {bp['D39'].value}; {bp['E39'].value}")
# ліміти п. 7
limits = {"Ліцензійне програмне": 0.25, "Вебсайт": 0.10, "Налаштування": 0.20, "Маркетинг": 0.10, "Фірмовий стиль": 0.50, "Оренда": 0.25}
total = 0
for r in range(82, 89):
    name, amt = bp[f"B{r}"].value, bp[f"E{r}"].value or 0
    total += amt
    for key, lim in limits.items():
        if name and name.startswith(key):
            check(amt <= lim * grant, f"{name}: {amt:,.0f} ≤ {lim:.0%} × грант ({lim*grant:,.0f})")
check(total == grant, f"сума інвестплану = грант ({total:,.0f})")
check(bp["E96"].value == grant + M.OWN_TOTAL, f"загальний бюджет E96 = грант + власні ({bp['E96'].value:,.0f})")
# фінанси
col = lambda m: get_column_letter(4 + m)
taxes_total = pl["AN71"].value
check(abs(taxes_total - M.summary()["taxes"]) < 1, f"накопичені податки за 36 міс = модель ({taxes_total:,.2f})")
check(taxes_total >= grant, f"податки, збори та ЄСВ за строк проєкту ({taxes_total:,.0f}) ≥ грант ({grant:,.0f})")
check(pl["AN50"].value == 0, "залишок зобов'язання за грантом на М36 = 0")
cover = next((m for m in range(1, 37) if pl[f"{col(m)}50"].value == 0), None)
print(f"info покриття гранту податками — з М{cover}")
balances = [pl[f"{col(m)}78"].value for m in range(1, 37)]
check(min(balances) > 0, f"мінімальний залишок коштів {min(balances):,.2f} (М{balances.index(min(balances))+1}) > 0")
payroll = sum(pl[f"{col(m)}49"].value for m in range(1, 37))
check(payroll + taxes_total >= grant, f"зарплата ({payroll:,.0f}) + податки ({taxes_total:,.0f}) ≥ грант — вимога абз. 8 п. 19 Порядку")
check(bp["B99"].value >= 1 and bp["B109"].value == bp["B99"].value, f"нових робочих місць: {bp['B99'].value}")
check(pl["BM27"].value == pl["BM14"].value + pl["BM15"].value + pl["BM16"].value + pl["BM17"].value, "BM27 = сума надходжень")
check(abs(pl["BM74"].value - (pl["BM27"].value - pl["BM72"].value)) < 0.01 and abs(pl["BM78"].value - pl["BM74"].value) < 0.01, "BM74/BM78 узгоджені")
# текстові цифри ↔ таблиця
def nums(s): return [int(x.replace(" ", "")) for x in re.findall(r"\d[\d ]{4,}", s or "")]
rev_total = pl["BM16"].value + pl["BM17"].value
check(int(round(rev_total)) in nums(bp["B62"].value), f"виручка у тексті B62 = таблиця ({rev_total:,.0f})")
check(int(round(taxes_total)) in nums(bp["B111"].value), f"податки у тексті B111 = таблиця ({taxes_total:,.0f})")
check(int(round(pl["BM78"].value)) in nums(bp["B65"].value), f"залишок коштів у тексті B65 = таблиця ({pl['BM78'].value:,.0f})")
# формули шаблону збережені (крім свідомо змінених)
allowed = {"BM27","BM46","BM74"} | {f"{col(m)}{r}" for m in range(1,61) for r in (60,61,66)}
changed = []
for name in ("Бізнес план", "Прибутки та збитки"):
    a, b = wbt[name], wbf[name]
    for row in a.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("=") and b[c.coordinate].value != c.value and c.coordinate not in allowed:
                changed.append((name, c.coordinate))
check(not changed, f"формули шаблону не перезаписані (змінено лише дозволені): {changed[:10]}")
# відсутність плейсхолдерів, які треба заповнити
ph = []
for ws in (bp, pl):
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and "[" in c.value and "]" in c.value:
                ph.append((ws.title, c.coordinate, re.findall(r"\[[^\]]+\]", c.value)))
print("info поля для уточнення заявницею:", ph)
print("\nРЕЗУЛЬТАТ:", "усі перевірки пройдено" if ok else "є помилки")
sys.exit(0 if ok else 1)
